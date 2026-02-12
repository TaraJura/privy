from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document

from .detector import BaseDetector
from .docx_engine import apply_replacements_to_paragraph, iter_document_paragraphs, paragraph_text
from .mapping_store import MappingData, read_mapping, write_mapping
from .types import SpanReplacement, VALID_ENTITY_TYPES


class AnonymizationError(RuntimeError):
    pass


# Common legal document role words used in party descriptors (e.g. "THE CONSULTANT").
# These are structural labels, not actual entity names, and should not be anonymized.
_LEGAL_ROLE_WORDS = frozenset({
    "AGENT", "ASSIGNEE", "ASSIGNOR", "BENEFICIARY", "BORROWER",
    "BUYER", "CLIENT", "COMPANY", "CONSULTANT", "CONTRACTOR",
    "CUSTOMER", "DISTRIBUTOR", "EMPLOYEE", "EMPLOYER", "EXECUTOR",
    "GUARANTOR", "INVESTOR", "LANDLORD", "LENDER", "LESSEE",
    "LESSOR", "LICENSEE", "LICENSOR", "PARTNER", "PARTNERS",
    "PARTIES", "PARTY", "PRINCIPAL", "PROVIDER", "RECIPIENT",
    "SELLER", "SUBCONTRACTOR", "SUPPLIER", "TENANT", "TRUSTEE",
    "VENDOR",
})


@dataclass
class ProcessingReport:
    paragraphs_scanned: int
    replacements_applied: int
    entities_detected: int

    def to_dict(self) -> dict[str, int]:
        return {
            "paragraphs_scanned": self.paragraphs_scanned,
            "replacements_applied": self.replacements_applied,
            "entities_detected": self.entities_detected,
        }


def anonymize_docx(
    input_path: Path,
    output_path: Path,
    map_path: Path,
    detector: BaseDetector,
    entity_types: Iterable[str],
    min_confidence: float = 0.5,
    report_path: Path | None = None,
) -> ProcessingReport:
    normalized_entity_types = _normalize_entity_types(entity_types)

    doc = Document(str(input_path))
    mapping = MappingData.create_empty()
    reverse_index: dict[tuple[str, str], str] = {}
    counters = {entity_type: 0 for entity_type in normalized_entity_types}

    paragraphs_scanned = 0
    replacements_applied = 0
    entities_detected = 0

    for paragraph_ref in iter_document_paragraphs(doc):
        text = paragraph_text(paragraph_ref.paragraph)
        if not text.strip():
            continue

        paragraphs_scanned += 1
        raw_entities = detector.detect(text)
        entities = _select_entities(raw_entities, normalized_entity_types, min_confidence)
        if not entities:
            continue

        entities_detected += len(entities)

        replacements: list[SpanReplacement] = []
        for entity in entities:
            original = text[entity.start : entity.end]
            key = (entity.label, original)
            placeholder = reverse_index.get(key)
            if placeholder is None:
                counters[entity.label] += 1
                placeholder = f"{entity.label}_{counters[entity.label]:03d}"
                reverse_index[key] = placeholder
                mapping.placeholders[placeholder] = {
                    "label": entity.label,
                    "original": original,
                }
            replacements.append(
                SpanReplacement(start=entity.start, end=entity.end, replacement=placeholder)
            )

        replacements_applied += apply_replacements_to_paragraph(paragraph_ref.paragraph, replacements)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    write_mapping(path=map_path, mapping=mapping)

    report = ProcessingReport(
        paragraphs_scanned=paragraphs_scanned,
        replacements_applied=replacements_applied,
        entities_detected=entities_detected,
    )

    if report_path is not None:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report.to_dict(), indent=2), encoding="utf-8")

    return report


def deanonymize_docx(
    input_path: Path,
    output_path: Path,
    map_path: Path,
    report_path: Path | None = None,
) -> ProcessingReport:
    mapping = read_mapping(path=map_path)
    if not mapping.placeholders:
        raise AnonymizationError("Mapping has no placeholders. Nothing to deanonymize.")

    doc = Document(str(input_path))

    paragraphs_scanned = 0
    replacements_applied = 0
    entities_detected = 0

    for paragraph_ref in iter_document_paragraphs(doc):
        text = paragraph_text(paragraph_ref.paragraph)
        if not text.strip():
            continue

        paragraphs_scanned += 1
        replacements = _placeholder_replacements(text, mapping)
        if not replacements:
            continue

        entities_detected += len(replacements)
        replacements_applied += apply_replacements_to_paragraph(paragraph_ref.paragraph, replacements)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    report = ProcessingReport(
        paragraphs_scanned=paragraphs_scanned,
        replacements_applied=replacements_applied,
        entities_detected=entities_detected,
    )

    if report_path is not None:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report.to_dict(), indent=2), encoding="utf-8")

    return report


def _normalize_entity_types(entity_types: Iterable[str]) -> set[str]:
    normalized = set()
    for entity_type in entity_types:
        value = entity_type.strip().upper()
        if value not in VALID_ENTITY_TYPES:
            raise AnonymizationError(
                f"Unsupported entity type: {entity_type}. "
                f"Valid values: {', '.join(sorted(VALID_ENTITY_TYPES))}."
            )
        normalized.add(value)

    if not normalized:
        raise AnonymizationError("At least one entity type must be provided.")

    return normalized


def _select_entities(
    entities: list[EntitySpan],
    entity_types: set[str],
    min_confidence: float,
) -> list[EntitySpan]:
    candidates = [
        entity
        for entity in entities
        if entity.label in entity_types
        and entity.confidence >= min_confidence
        and entity.end > entity.start
        and not _is_legal_role_label(entity.text)
    ]

    if not candidates:
        return []

    ranked = sorted(
        candidates,
        key=lambda e: (-e.confidence, -(e.end - e.start), e.start),
    )

    selected: list[EntitySpan] = []
    for candidate in ranked:
        if any(_overlaps(candidate.start, candidate.end, current.start, current.end) for current in selected):
            continue
        selected.append(candidate)

    return sorted(selected, key=lambda e: e.start)


def _placeholder_replacements(text: str, mapping: MappingData) -> list[SpanReplacement]:
    matches: list[tuple[int, int, str]] = []
    sorted_placeholders = sorted(mapping.placeholders.items(), key=lambda item: len(item[0]), reverse=True)

    for placeholder, metadata in sorted_placeholders:
        original = metadata.get("original")
        if not original:
            continue

        start = 0
        while True:
            idx = text.find(placeholder, start)
            if idx == -1:
                break
            matches.append((idx, idx + len(placeholder), original))
            start = idx + len(placeholder)

    if not matches:
        return []

    matches.sort(key=lambda item: (item[0], -(item[1] - item[0])))

    replacements: list[SpanReplacement] = []
    current_end = -1
    for start, end, original in matches:
        if start < current_end:
            continue
        replacements.append(SpanReplacement(start=start, end=end, replacement=original))
        current_end = end

    return replacements


def _is_legal_role_label(text: str) -> bool:
    """Return True if *text* is a legal document role label like 'THE CONSULTANT'."""
    upper = text.strip().upper()
    if not upper.startswith("THE "):
        return False
    return upper[4:].strip() in _LEGAL_ROLE_WORDS


def _overlaps(start_a: int, end_a: int, start_b: int, end_b: int) -> bool:
    return start_a < end_b and start_b < end_a
